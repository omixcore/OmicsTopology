#!/usr/bin/env python3
"""Build PLOS Computational Biology manuscript, SI, and vector figure assets from versioned sources."""
from pathlib import Path
import base64, gzip, io, re, shutil, subprocess
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import cairosvg
from pypdf import PdfWriter
from PIL import Image
ROOT=Path(__file__).resolve().parent
ASSET_DIR=ROOT/'asset_sources'; MAIN_FIG=ROOT/'figures'; SUPP_FIG=ROOT/'supplementary_figures'

def expand_sources():
    MAIN_FIG.mkdir(parents=True,exist_ok=True); SUPP_FIG.mkdir(parents=True,exist_ok=True)
    for src in ASSET_DIR.glob('*.gz.b64'):
        raw=base64.b64decode(src.read_text(encoding='ascii').strip())
        name=src.name[:-len('.gz.b64')]
        if name.endswith('.md'): dest=ROOT/name
        elif name.endswith('.svg'): dest=(MAIN_FIG if name.startswith('Fig') else SUPP_FIG)/name
        else: continue
        dest.parent.mkdir(parents=True,exist_ok=True)
        with gzip.GzipFile(fileobj=io.BytesIO(raw),mode='rb') as fin, open(dest,'wb') as fout: shutil.copyfileobj(fin,fout)

def set_page_number(section):
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.clear()
    r=p.add_run('Page '); r.font.name='Times New Roman'; r.font.size=Pt(10)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); p._p.append(field)

def set_line_numbers(section):
    sect=section._sectPr; old=sect.find(qn('w:lnNumType'))
    if old is not None: sect.remove(old)
    ln=OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'),'1'); ln.set(qn('w:restart'),'continuous'); ln.set(qn('w:distance'),'360'); ln.set(qn('w:lineNumber'),'1'); sect.append(ln)

def configure_doc(doc,line_numbers=False):
    for sec in doc.sections:
        sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); set_page_number(sec)
        if line_numbers: set_line_numbers(sec)
    for name in ('Normal','Body Text','Caption'):
        try:
            st=doc.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(12); st.paragraph_format.line_spacing=2.0; st.paragraph_format.space_after=Pt(0)
        except KeyError: pass
    for name in ('Heading 1','Heading 2','Heading 3'):
        try:
            st=doc.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(12); st.font.bold=True; st.paragraph_format.line_spacing=2.0; st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(0)
        except KeyError: pass

def markdown_blocks(path):
    lines=path.read_text(encoding='utf-8').splitlines(); i=0
    while i<len(lines):
        line=lines[i].rstrip()
        if not line.strip(): i+=1; continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i+=1
            rows=[]
            for row in block:
                cells=[c.strip() for c in row.strip('|').split('|')]
                if cells and all(re.fullmatch(r':?-{3,}:?',c.replace(' ','')) for c in cells): continue
                rows.append(cells)
            if rows: yield ('table',rows)
            continue
        if line.startswith('!['):
            m=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
            if m: yield ('image',(m.group(1),m.group(2)))
            i+=1; continue
        if line.startswith('#'):
            n=len(line)-len(line.lstrip('#')); yield ('heading',(n,line[n:].strip())); i+=1; continue
        para=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith('#') and not lines[i].startswith('|') and not lines[i].startswith('!['): para.append(lines[i].strip()); i+=1
        yield ('paragraph',' '.join(para))

def add_table(doc,rows):
    width=max(len(r) for r in rows); table=doc.add_table(rows=1,cols=width); table.style='Table Grid'
    for j,val in enumerate(rows[0]): table.rows[0].cells[j].text=val
    for row in rows[1:]:
        cells=table.add_row().cells
        for j in range(width): cells[j].text=row[j] if j<len(row) else ''
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(0)
                for run in p.runs: run.font.name='Times New Roman'; run.font.size=Pt(8 if width>=5 else 9); run.bold=(ri==0)

def build_docx(md_path,out_path,manuscript=True):
    doc=Document(); configure_doc(doc,line_numbers=manuscript)
    for kind,value in markdown_blocks(md_path):
        if kind=='heading':
            level,text=value
            if manuscript and text=='Abstract':
                p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.page_break_before=True; p.add_run(text)
            elif level==1:
                p=doc.add_paragraph(style='Title' if manuscript else 'Heading 1'); p.add_run(text)
                if manuscript:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(16); r.bold=True
            else: doc.add_paragraph(text,style='Heading 1' if level==2 else 'Heading 2')
        elif kind=='paragraph':
            p=doc.add_paragraph(value,style='Normal')
            if manuscript and value.startswith('Xiaofeng Song'): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if manuscript and value.startswith('Short title:'):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name='Times New Roman'; run.font.size=Pt(10); run.italic=True
            if manuscript and (value.startswith('1 North Henan') or value.startswith('2 Key Laboratory') or value.startswith('† These authors')): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=='image':
            alt,rel=value; svg=(ROOT/rel).resolve(); png=svg.with_suffix('.png')
            if not png.exists(): cairosvg.svg2png(url=str(svg),write_to=str(png))
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(png),width=Inches(6.25))
        elif kind=='table': add_table(doc,value)
    doc.save(out_path)

def make_pdf(docx_path):
    out=docx_path.with_suffix('.pdf')
    if out.exists(): out.unlink()
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(ROOT),str(docx_path)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)

def merge_pdfs(paths,out):
    writer=PdfWriter()
    for path in paths: writer.append(str(path))
    with open(out,'wb') as f: writer.write(f)

def main():
    expand_sources()
    for folder in (MAIN_FIG,SUPP_FIG):
        for svg in sorted(folder.glob('*.svg')):
            cairosvg.svg2png(url=str(svg),write_to=str(svg.with_suffix('.png')))
            cairosvg.svg2pdf(url=str(svg),write_to=str(svg.with_suffix('.pdf')))
            if folder == MAIN_FIG:
                # PLOS figure-upload guidance requests TIFF/EPS at 300-600 ppi.
                # Keep the SVG as the editable source and write a 600-ppi RGB TIFF.
                temp_png = svg.with_name(svg.stem + '.submission-600dpi.png')
                cairosvg.svg2png(url=str(svg),write_to=str(temp_png),output_width=3600)
                with Image.open(temp_png) as image_file:
                    image_file.convert('RGB').save(svg.with_suffix('.tif'),format='TIFF',dpi=(600,600),compression='tiff_lzw')
                temp_png.unlink(missing_ok=True)
    m=ROOT/'OmicsTopology_PLOS_Comp_Biol_manuscript.md'; s=ROOT/'OmicsTopology_PLOS_Comp_Biol_supporting_information.md'
    md=ROOT/'OmicsTopology_PLOS_Comp_Biol_manuscript.docx'; sd=ROOT/'OmicsTopology_PLOS_Comp_Biol_supporting_information.docx'
    build_docx(m,md,True); build_docx(s,sd,False); make_pdf(md); make_pdf(sd)
    merge_pdfs([MAIN_FIG/f'Fig{i}.pdf' for i in range(1,7)],ROOT/'Main_figures.pdf')
    merge_pdfs([SUPP_FIG/f'S{i}_Fig.pdf' for i in range(1,12)],ROOT/'Supplementary_figures.pdf')
    print('Built PLOS manuscript, supporting information, and 17 vector figure sets.')
if __name__=='__main__': main()
